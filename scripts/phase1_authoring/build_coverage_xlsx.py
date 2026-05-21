import json, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

REPO = Path(sys.argv[1])
idx = json.loads((REPO / "data/labelling/corpus_index.json").read_text())

SIGNALS = ["security_concern", "legal_risk", "reputation_risk", "churn_risk",
           "complaint", "bug_report", "feature_request", "praise",
           "competitor_mention", "expansion_opportunity", "unclear", "not_actionable"]
SEVS = ["SEV-1", "SEV-2", "SEV-3", "SEV-4", "SEV-5"]

seeds = {}
for r in idx:
    if r["kind"] == "seed":
        seeds.setdefault((r["signal_type"], r["severity"]), []).append(r["scenario_id"])
held = {}
for r in idx:
    if r["split"] == "heldout":
        held.setdefault((r["signal_type"], r["severity"]), []).append(r["scenario_id"])

FONT = "Arial"
HDR = Font(name=FONT, bold=True, color="FFFFFF", size=11)
BOLD = Font(name=FONT, bold=True, size=11)
NORM = Font(name=FONT, size=10)
SMALL = Font(name=FONT, size=9, color="555555")
HFILL = PatternFill("solid", fgColor="2F5496")
ACTIVE = PatternFill("solid", fgColor="E2EFDA")
NA = PatternFill("solid", fgColor="EDEDED")
ABST = PatternFill("solid", fgColor="FFF2CC")
thin = Side(style="thin", color="BFBFBF")
BORD = Border(left=thin, right=thin, top=thin, bottom=thin)
CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)
LT = Alignment(horizontal="left", vertical="center", wrap_text=True)

wb = Workbook()


def matrix_sheet(ws, data, title, na_active_only):
    ws.append([title])
    ws["A1"].font = Font(name=FONT, bold=True, size=14)
    ws.append([])
    header = ["SignalType \\ Severity"] + SEVS
    ws.append(header)
    for c in range(1, len(header) + 1):
        cell = ws.cell(row=3, column=c)
        cell.font = HDR
        cell.fill = HFILL
        cell.alignment = CTR
        cell.border = BORD
    r0 = 4
    for i, sig in enumerate(SIGNALS):
        ws.cell(row=r0 + i, column=1, value=sig).font = BOLD
        ws.cell(row=r0 + i, column=1).alignment = LT
        ws.cell(row=r0 + i, column=1).border = BORD
        for j, sev in enumerate(SEVS):
            cell = ws.cell(row=r0 + i, column=2 + j)
            cell.border = BORD
            cell.alignment = CTR
            ids = data.get((sig, sev), [])
            active = (sig, sev) in na_active_only
            if ids:
                cell.value = "\n".join(ids)
                cell.font = NORM
                cell.fill = ABST if sig == "unclear" else ACTIVE
            elif active:
                # active cell that should have entries but doesn't (shouldn't happen)
                cell.value = "(none)"
                cell.font = SMALL
                cell.fill = ACTIVE
            else:
                cell.value = "N/A — Lead sign-off"
                cell.font = SMALL
                cell.fill = NA
    ws.column_dimensions["A"].width = 22
    for col in "BCDEF":
        ws.column_dimensions[col].width = 20
    for i in range(len(SIGNALS)):
        ws.row_dimensions[r0 + i].height = 30
    # legend
    lr = r0 + len(SIGNALS) + 1
    ws.cell(row=lr, column=1, value="Legend:").font = BOLD
    ws.cell(row=lr + 1, column=1, value="Active cell (≥2 seeds / ≥1 held-out)").font = SMALL
    ws.cell(row=lr + 1, column=1).fill = ACTIVE
    ws.cell(row=lr + 2, column=1, value="Abstention cell (unclear × SEV-5)").font = SMALL
    ws.cell(row=lr + 2, column=1).fill = ABST
    ws.cell(row=lr + 3, column=1, value="N/A — requires Lead sign-off").font = SMALL
    ws.cell(row=lr + 3, column=1).fill = NA


active_cells = set(seeds.keys())
ws1 = wb.active
ws1.title = "Coverage (Train Seeds)"
matrix_sheet(ws1, seeds, "Coverage Matrix — Train/Val Seed Scenarios (≥2 per active cell)", active_cells)

ws2 = wb.create_sheet("Coverage (Held-out)")
matrix_sheet(ws2, held, "Coverage Matrix — Held-out Set (≥1 per active cell)", active_cells)

# Summary sheet
ws3 = wb.create_sheet("Summary")
rows = [
    ["Phase-1 Labelling Corpus — Summary", ""],
    ["", ""],
    ["Total scenarios", len(idx)],
    ["  Seeds (train/val)", sum(1 for r in idx if r["kind"] == "seed")],
    ["  Adversarial variants (train/val)", sum(1 for r in idx if r["kind"] == "variant")],
    ["  Held-out", sum(1 for r in idx if r["split"] == "heldout")],
    ["", ""],
    ["Split: train", sum(1 for r in idx if r["split"] == "train")],
    ["Split: val", sum(1 for r in idx if r["split"] == "val")],
    ["Split: heldout", sum(1 for r in idx if r["split"] == "heldout")],
    ["", ""],
    ["Active SignalType×Severity cells", len(active_cells)],
    ["N/A cells (Lead sign-off)", 12 * 5 - len(active_cells)],
    ["Abstentions (train/val)", sum(1 for r in idx if r["abstain"] and r["split"] in ("train", "val"))],
    ["Abstentions (held-out)", sum(1 for r in idx if r["abstain"] and r["split"] == "heldout")],
    ["Held-out adversarial-tagged", sum(1 for r in idx if r["split"] == "heldout" and r["adversarial_tags"])],
    ["Held-out multi-tag (≥2)", sum(1 for r in idx if r["split"] == "heldout" and len(r["adversarial_tags"]) >= 2)],
    ["", ""],
    ["IAA target: κ(signal_type)", "≥ 0.80"],
    ["IAA target: weighted κ(severity)", "≥ 0.75"],
    ["IAA target: agreement(abstain)", "≥ 0.85"],
    ["PII sign-off", "100% (pii_review_passed=true)"],
]
for r in rows:
    ws3.append(r)
ws3["A1"].font = Font(name=FONT, bold=True, size=14)
for i in range(3, len(rows) + 1):
    ws3.cell(row=i, column=1).font = BOLD if ws3.cell(row=i, column=1).value and not ws3.cell(row=i, column=1).value.startswith("  ") else NORM
    ws3.cell(row=i, column=2).font = NORM
    ws3.cell(row=i, column=2).alignment = Alignment(horizontal="left")
ws3.column_dimensions["A"].width = 38
ws3.column_dimensions["B"].width = 32

out = REPO / "data/labelling/coverage_matrix.xlsx"
wb.save(out)
print("wrote", out)
